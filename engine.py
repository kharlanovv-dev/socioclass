import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_excel_report(school, cls_name, expert, parsed_data, file_path, threshold_star=3, threshold_rej=0, country="РБ", calc_indices=False):
    count = len(parsed_data)
    wb = openpyxl.Workbook()
    
    font_title = Font(name="Times New Roman", size=12, bold=True)
    font_header = Font(name="Times New Roman", size=12, bold=True)
    font_main = Font(name="Times New Roman", size=12)
    font_italic = Font(name="Times New Roman", size=12, italic=True)
    font_alert = Font(name="Times New Roman", size=11, italic=True, color="FF0000")
    
    fill_diag = PatternFill(fill_type="solid", fgColor="D3D3D3")
    fill_p = PatternFill(fill_type="solid", fgColor="C6EFCE")
    fill_m = PatternFill(fill_type="solid", fgColor="FFC7CE")
    
    border_thin = Side(border_style="thin", color="000000")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # --- ЛИСТ 1: МАТРИЦА ---
    ws1 = wb.active
    ws1.title = "Матрица выборов"
    ws1.sheet_view.showGridLines = True
    
    ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0
    ws1.sheet_properties.pageSetUpPr.fitToPage = True
    
    school_prefix = "ГУО / УО" if country == "РБ" else "МБОУ СОШ / Школа"
    ws1.cell(row=1, column=1, value=f"Социометрическая матрица ({country}). {school_prefix}: {school}, класс {cls_name}").font = font_title
    
    matrix = [['' for _ in range(count)] for _ in range(count)]
    for i in range(1, count + 1):
        for p in parsed_data[i]['plus']:
            if 1 <= p <= count: matrix[i-1][p-1] = '+'
        for m in parsed_data[i]['minus']:
            if 1 <= m <= count: matrix[i-1][m-1] = '-'
            
    for i in range(1, count + 1):
        c_row = ws1.cell(row=3, column=i+1, value=i)
        c_row.font = font_header; c_row.border = box_border; c_row.alignment = align_center
        c_col = ws1.cell(row=i+3, column=1, value=f"Ученик {i}")
        c_col.font = font_main; c_col.border = box_border
        
    headers_end = ["+", "±", "++", "--"]
    for idx, h in enumerate(headers_end):
        c = ws1.cell(row=3, column=count + 2 + idx, value=h)
        c.font = font_header; c.border = box_border; c.alignment = align_center
        
    stats = {i: {'p': 0, 'm': 0, 'pp': 0, 'mm': 0} for i in range(1, count + 1)}
    mut_p, mut_m = [], []
    
    for i in range(1, count + 1):
        for j in range(1, count + 1):
            cell = ws1.cell(row=i+3, column=j+1)
            cell.border = box_border; cell.alignment = align_center; cell.font = font_main
            if i == j:
                cell.fill = fill_diag; cell.value = "X"
                continue
            val = matrix[i-1][j-1]
            rev_val = matrix[j-1][i-1]
            cell.value = val
            if val == '+':
                stats[j]['p'] += 1
                if rev_val == '+':
                    cell.fill = fill_p
                    stats[i]['pp'] += 0.5
                    if i < j: mut_p.append((i, j))
            elif val == '-':
                stats[j]['m'] += 1
                if rev_val == '-':
                    cell.fill = fill_m
                    stats[i]['mm'] += 0.5
                    if i < j: mut_m.append((i, j))
                    
    for i in range(1, count + 1):
        for idx, key in enumerate(['p', 'm', 'pp', 'mm']):
            val = stats[i][key] if key in ['p', 'm'] else int(stats[i][key] * 2)
            c = ws1.cell(row=i+3, column=count + 2 + idx, value=val)
            c.border = box_border; c.font = font_main; c.alignment = align_center
            
    ws1.column_dimensions['A'].width = 18
    for col in range(2, count + 6):
        ws1.column_dimensions[get_column_letter(col)].width = 3.5

    # --- ЛИСТ 2: АНАЛИТИЧЕСКИЙ ОТЧЕТ ---
    ws2 = wb.create_sheet(title="Аналитический отчет")
    ws2.sheet_view.showGridLines = True
    ws2.page_setup.orientation = ws2.ORIENTATION_PORTRAIT
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.column_dimensions['A'].width = 85
    
    ws2.cell(row=2, column=1, value="РЕЗУЛЬТАТЫ СОЦИОМЕТРИЧЕСКОГО ИССЛЕДОВАНИЯ").font = font_title
    ws2.cell(row=3, column=1, value=f"{school_prefix}: {school} | Класс: {cls_name}").font = font_italic
    ws2.cell(row=4, column=1, value="⚠️ Если отчет пуст, нажмите «Разрешить редактирование» вверху окна Excel.").font = font_alert
    
    stars_raw, accepted_raw, rejected_raw, isolated_idx = [], [], [], []
    for i in range(1, count + 1):
        p = stats[i]['p']
        m = stats[i]['m']
        if p == 0 and m == 0: isolated_idx.append(i)
        elif p >= threshold_star: stars_raw.append((i, p))
        elif p <= threshold_rej: rejected_raw.append((i, p))
        else: accepted_raw.append((i, p))
        
    stars_idx = sorted(stars_raw, key=lambda x: x[1], reverse=True)
    accepted_idx = sorted(accepted_raw, key=lambda x: x[1], reverse=True)
    rejected_idx = sorted(rejected_raw, key=lambda x: x[1], reverse=True)
        
    def write_dynamic_section(ws, row, title, idx_list, mode="single"):
        ws.cell(row=row, column=1, value=title).font = font_header
        if not idx_list:
            ws.cell(row=row+1, column=1, value="не обнаружено").font = font_main
            return row + 3
        formula_parts = []
        if mode == "single":
            for idx in idx_list: formula_parts.append(f"'Матрица выборов'!A{idx+3}")
        elif mode == "with_score":
            for idx, score in idx_list:
                formula_parts.append(f"'Матрица выборов'!A{idx+3}")
                formula_parts.append(f'" (выборов: {score})"')
        elif mode == "pairs":
            for idx1, idx2 in idx_list:
                formula_parts.append(f"'Матрица выборов'!A{idx1+3}")
                formula_parts.append(f'" + "')
                formula_parts.append(f"'Матрица выборов'!A{idx2+3}")
        final_formula_args = []
        for i, part in enumerate(formula_parts):
            final_formula_args.append(part)
            if mode == "with_score" and i % 2 == 1 and i < len(formula_parts) - 1: final_formula_args.append('", "')
            elif mode == "single" and i < len(formula_parts) - 1: final_formula_args.append('", "')
            elif mode == "pairs" and i % 3 == 2 and i < len(formula_parts) - 1: final_formula_args.append('"; "')
            
        full_formula = f"=CONCATENATE({', '.join(final_formula_args)})"
        c = ws.cell(row=row+1, column=1, value=full_formula)
        c.font = font_main; c.alignment = align_left_wrap
        ws.row_dimensions[row+1].height = max(20, (len(idx_list) * 25 // 80 + 1) * 18)
        return row + 3
        
    r = 7
    r = write_dynamic_section(ws2, r, f"1. Группа «Звезды» (получили ПЛЮСОВ >= {threshold_star}):", stars_idx, "with_score")
    r = write_dynamic_section(ws2, r, "2. Группа «Принятые» (средний статус):", accepted_idx, "with_score")
    r = write_dynamic_section(ws2, r, f"3. Группа «Отвергаемые» (получили ПЛЮСОВ <= {threshold_rej}):", rejected_idx, "with_score")
    r = write_dynamic_section(ws2, r, "4. Группа «Изолированные» (0 плюсов и 0 минусов):", isolated_idx, "single")
    r = write_dynamic_section(ws2, r, f"5. Взаимные положительные выборы (Всего: {len(mut_p)}):", mut_p, "pairs")
    r = write_dynamic_section(ws2, r, f"6. Взаимные отрицательные выборы (Всего: {len(mut_m)}):", mut_m, "pairs")
    
    if calc_indices:
        ws2.cell(row=r, column=1, value="7. ДОПОЛНИТЕЛЬНЫЕ СОЦИОМЕТРИЧЕСКИЕ ИНДЕКСЫ ГРУППЫ:").font = font_header
        r += 1
        total_good = len(stars_idx) + len(accepted_idx)
        ibv_percent = (total_good / count) * 100 if count > 0 else 0
        lvl = "Высокий уровень" if ibv_percent >= 75 else ("Средний уровень" if ibv_percent >= 50 else "Низкий уровень")
        ws2.cell(row=r, column=1, value=f"• Индекс благополучия взаимоотношений (ИБВ) = {ibv_percent:.1f}% ──> Status: {lvl}").font = font_main; r += 1
        
        max_possible_mut = (count * (count - 1)) / 2 if count > 1 else 1
        igs_val = len(mut_p) / max_possible_mut if max_possible_mut > 0 else 0
        coh = "Высокая сплоченность" if igs_val >= 0.4 else ("Средняя сплоченность" if igs_val >= 0.2 else "Низкая сплоченность")
        ws2.cell(row=r, column=1, value=f"• Индекс групповой сплоченности (ИГС) = {igs_val:.2f} ──> Характеристика: {coh}").font = font_main; r += 2
    
    ws2.cell(row=r+1, column=1, value=f"Диагностику провел: {expert}").font = font_header
    wb.save(file_path)
